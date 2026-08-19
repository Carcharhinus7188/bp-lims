# -*- coding: utf-8 -*-
"""受控原始记录 Word 引擎 — 填入实验数据 + 电子签名

Ported from bp-lims-reference/template_record_engine.py + record_word_engine.py
"""
from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

CHINA_TZ = timezone(timedelta(hours=8))

BLACK = RGBColor(0, 0, 0)
RED = RGBColor(255, 0, 0)

BLANK_RE = re.compile(r"_{2,}|＿{2,}|…{2,}")
SPACE_RE = re.compile(r"\s+")
DATE_RE = re.compile(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})")
TIME_RE = re.compile(r"(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?")

# ── 工具函数 ──────────────────────────────────────────────

def _clean(value: Any) -> str:
    return SPACE_RE.sub(" ", str(value or "").replace("\xa0", " ")).strip()


def _china_date_str(value: Any) -> str:
    """将时间值统一转为中国时区日期字符串 YYYY-MM-DD。

    asyncpg 返回的 timestamptz 是 UTC 的 aware datetime，直接 str()[:10] 会在
    00:00–08:00（北京时间）取到前一天，这里统一转北京时间再取日期。
    """
    if not value:
        return ""
    if isinstance(value, str):
        return value[:10]
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(CHINA_TZ)
        return value.strftime("%Y-%m-%d")
    return str(value)[:10]


def _fill_date(original: str, raw: str) -> str:
    match = DATE_RE.search(str(raw or ""))
    if not match:
        return original
    year, month, day = match.groups()
    groups = list(BLANK_RE.finditer(original))
    if len(groups) < 3:
        return BLANK_RE.sub(str(raw), original, count=1)
    replacements = [year, f"{int(month):02d}", f"{int(day):02d}"]
    output = original
    for marker, replacement in zip(groups, replacements):
        start, end = marker.start(), marker.end()
        output = output[:start] + str(replacement) + " " * (end - start - len(str(replacement))) + output[end:]
    return output


def _fill_time(original: str, raw: str) -> str:
    match = TIME_RE.search(str(raw or ""))
    if not match:
        return original
    hour, minute, second = match.groups()
    replacements = [f"{int(hour):02d}", f"{int(minute):02d}"]
    if second is not None:
        replacements.append(f"{int(second):02d}")
    groups = list(BLANK_RE.finditer(original))
    output = original
    for marker, replacement in zip(groups, replacements):
        start, end = marker.start(), marker.end()
        output = output[:start] + str(replacement) + " " * (end - start - len(str(replacement))) + output[end:]
    return output


def _compose_cell_text(original: str, raw_value: Any) -> str:
    raw = str(raw_value or "").strip()
    if not original:
        return raw
    if not raw:
        return original
    if ("□" in original or "☐" in original) and ("☑" in raw or "□" in raw):
        return raw
    if "年" in original and "月" in original and "日" in original and BLANK_RE.search(original):
        return _fill_date(original, raw)
    if ":" in original and BLANK_RE.search(original) and TIME_RE.search(raw):
        return _fill_time(original, raw)
    if BLANK_RE.search(original):
        return BLANK_RE.sub(raw, original, count=1)
    return raw


_NORM_RE = re.compile(r"[\s：:：/\\、，,；;（）()\[\]【】\-—_]+")


def _norm_label(value: Any) -> str:
    return _NORM_RE.sub("", str(value or "")).lower()


def _select_checkbox(text: str, preferred: Any) -> str:
    """在 checkbox 组中选择指定选项（否定词感知，避免“不符合”误匹配“符合”）。"""
    normalized = _norm_label(preferred)
    if not normalized:
        return text.replace("☑", "□")
    options = [x.strip() for x in re.split(r"[□☐☑]", text)[1:] if x.strip()]
    result = text.replace("☑", "□")
    matched = False
    for option in options:
        clean_opt = re.sub(r"[_＿…]+.*$", "", option).strip(" ：:；;，,")
        if not clean_opt:
            continue
        opt_norm = _norm_label(clean_opt)
        if opt_norm.startswith("不") and normalized == opt_norm[1:]:
            continue
        if normalized.startswith("不") and opt_norm == normalized[1:]:
            continue
        if opt_norm and (opt_norm in normalized or normalized in opt_norm):
            result = re.sub(r"□\s*" + re.escape(clean_opt),
                            lambda m: "☑" + m.group(0)[1:], result, count=1)
            matched = True
    if not matched:
        for option in options:
            clean_opt = re.sub(r"[_＿…]+.*$", "", option).strip(" ：:；;，,")
            if _norm_label(clean_opt) in {"其他", "其它"}:
                result = re.sub(r"□\s*" + re.escape(clean_opt),
                                lambda m: "☑" + m.group(0)[1:], result, count=1)
                break
    return result


def _select_checkbox_fill(original: str, value: Any) -> str:
    """勾选匹配项（无匹配则「其他」），并把剩余空白回填为实际值。

    对齐参考仓库 business_record_engine：样品规格/材料工艺等头部 checkbox
    在选定选项后，还要把选项后附的空白（___/＿/…）填入实际值（如规格型号/材料）。
    """
    selected = _select_checkbox(original, value)
    if BLANK_RE.search(selected):
        selected = BLANK_RE.sub(str(value or "").strip(), selected, count=1)
    return selected


def _contains_marker(text: str) -> bool:
    value = _clean(text)
    return not value or bool(BLANK_RE.search(value)) or "□" in value or "☐" in value


def _unique_row_cells(row) -> list[tuple[int, Any, str]]:
    result: list[tuple[int, Any, str]] = []
    seen = set()
    for col_index, cell in enumerate(row.cells):
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        result.append((col_index, cell, _clean(cell.text)))
    return result


def _table_headers(table) -> dict[int, str]:
    if not table.rows:
        return {}
    return {col_index: text for col_index, _, text in _unique_row_cells(table.rows[0])}


def _body_table_sections(doc: _Document) -> list[str]:
    sections: list[str] = []
    last_text = ""
    table_index = 0
    for child in doc._element.body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            text = _clean(Paragraph(child, doc._body).text)
            if text:
                last_text = text
        elif child.tag == qn("w:tbl"):
            sections.append(last_text or f"表{table_index + 1}")
            table_index += 1
    return sections


def _infer_input_type(original: str) -> str:
    if "□" in original or "☐" in original:
        return "checkbox"
    if "年" in original and "月" in original and "日" in original and BLANK_RE.search(original):
        return "date"
    if len(original) > 45:
        return "textarea"
    return "text"


def _row_has_explicit_marker(row) -> bool:
    """行内是否有勾选框/填空占位（□/☐/____ 等），用于区分表头行与数据行。"""
    seen = set()
    for _, cell, text in _unique_row_cells(row):
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        if text and (BLANK_RE.search(text) or "□" in text or "☐" in text):
            return True
    return False


def _secondary_col_header(table, row_index: int, col_index: int) -> str:
    """R0 表头为空/占位符时，向上合并连续表头行该列文本作为次表头。

    用于网格表里列名在非首行的情况（灰度测量「测量值1/2/3」、观察者信息
    「是否佩戴有色镜片」、厚度测量「固定端截面 P1」等多级表头），避免可填
    单元格都退化成行标签而重名。遇到含显式占位符（□/____）的行即停止，视为
    上一子表的边界。
    """
    parts: list[str] = []
    found_header = False
    for r in range(row_index - 1, 0, -1):  # 跳过 R0（常为整表标题）
        if r >= len(table.rows):
            continue
        row = table.rows[r]
        if _row_has_explicit_marker(row):
            # 数据/信息行（含勾选框或填空占位）：找到表头前继续向上跳，
            # 已找到表头后再遇到即视为上一子表边界而停止。
            if found_header:
                break
            continue
        if col_index < len(row.cells):
            text = _clean(row.cells[col_index].text)
            if text and not _contains_marker(text):
                parts.append(text)
                found_header = True
    parts.reverse()  # 自上而下
    return " ".join(parts)


_ROLE_TITLES = {
    "检测人员", "检验人员", "试验人员", "试验员", "检验员",
    "核验人员", "复核人员", "审核人员", "批准人员",
    "技术负责人", "授权签字人", "质量负责人", "批准人", "审核人",
    "记录人", "复核人", "编制人", "签发人", "采样人", "见证人", "校核人",
    "主检", "主检人", "检测人", "试验人", "检验人",
    "记录归档人", "归档人",
}


def _is_date_cell(text: str) -> bool:
    """是否纯日期填空格（如 ____年__月__日），不含人名/其他文本。"""
    stripped = re.sub(r"[_＿\s：:]+", "", text or "")
    return bool(stripped) and bool(re.fullmatch(r"[年月日0-9]+", stripped))


def _right_label(row_cells: list, col_index: int) -> str:
    """当前格右侧紧邻（下一个 col_index 更大）单元格的文本。

    用于角色行里区分「姓名 / 签字 / 日期」空白格：例如
    「检测人员 | 空白 | 签字：」中，第二个空白格右侧是「签字：」标签，
    应判为签字格而非第二个姓名格。
    """
    for c_col, _, c_text in row_cells:
        if c_col > col_index:
            return c_text.strip()
    return ""


def _nearest_role(row_cells: list, col_index: int) -> str:
    """当前格左侧最近的签字角色名（检测人员/核验人员/技术负责人…）。

    与 _role_row_name 不同：同一行可能横排多个角色
    （如「检测人员|签名|核验人员|签名」），此时按左侧最近的角色命名，
    而非行首最左角色，避免两个签名格都叫「检测人员 签字」。
    """
    role = ""
    for c_col, _, c_text in row_cells:
        if c_col >= col_index:
            break
        if not c_text or _contains_marker(c_text):
            continue
        base = re.sub(r"[（(][^）)]*[）)]", "", c_text.strip())
        if "日期" in base:
            continue
        parts = [p.strip() for p in re.split(r"[/、,，;；]", base) if p.strip()]
        if parts and any(p in _ROLE_TITLES for p in parts):
            role = c_text.strip()
    return role


def _is_first_blank_after_role(row_cells: list, col_index: int) -> bool:
    """当前格是否为左侧最近角色名之后第一个空白/标记格（即该角色的姓名格）。

    同一行可能横排多个角色（检测人员…姓名/日期 核验人员…姓名/日期），
    「第一个空白」应相对左侧最近的角色而言，而不是相对整行开头，
    否则核验人员的姓名格会被当成后续空白格而误判为日期。
    """
    last_role_col = None
    blank_after_role = False
    for c_col, _, c_text in row_cells:
        if c_col >= col_index:
            break
        base = re.sub(r"[（(][^）)]*[）)]", "", (c_text or "").strip())
        is_role = (
            bool(c_text)
            and not _contains_marker(c_text)
            and "日期" not in base
            and any(p.strip() in _ROLE_TITLES for p in re.split(r"[/、,，;；]", base) if p.strip())
        )
        if is_role:
            last_role_col = c_col
            blank_after_role = False
        elif last_role_col is not None and _contains_marker(c_text):
            blank_after_role = True
    return last_role_col is not None and not blank_after_role


def _role_row_name(row_cells: list) -> str:
    """判断是否为签字/角色行，返回行首角色名（检测人员/核验人员/技术负责人…）。

    附件归档表末尾的角色行与表格顶部列头（附件/记录名称、是否归档、备注）因单元格
    合并而结构不对齐，列头会错误覆盖日期/签字/结论格。非角色行返回空串。

    复合标签（如「检测人员/日期」「确认人员/日期」）含「日期」，属单个信息格而非
    角色行，予以排除，避免把信息行误当签字行重命名。
    """
    for _, _, t in row_cells:
        if t and not _contains_marker(t):
            leftmost = t.strip()
            break
    else:
        return ""
    base = re.sub(r"[（(][^）)]*[）)]", "", leftmost)  # 去掉「（如适用）」等括号注释
    if "日期" in base:
        return ""
    parts = [p.strip() for p in re.split(r"[/、,，;；]", base) if p.strip()]
    if not parts:
        return ""
    if any(p in _ROLE_TITLES for p in parts):
        return leftmost
    return ""


def _field_label(
    table_index: int, row_index: int, col_index: int,
    row_cells: list, headers: dict[int, str], template_text: str,
    table=None,
) -> tuple[str, str, str]:
    row_label = ""
    immediate_left = ""
    for candidate_col, _, candidate_text in row_cells:
        if candidate_col >= col_index:
            break
        # 跳过勾选框/填空占位等标记单元格（如「□ 是 □ 否」「____」），
        # 否则这类标记会被误当作行标签，导致同列的多个字段显示成同名重复项
        if candidate_text and not _contains_marker(candidate_text):
            row_label = candidate_text
            immediate_left = candidate_text
    col_header = headers.get(col_index, "")
    # 次表头：R0 表头无效时向上查纯表头行该列文本，补全网格表的列名
    if (not col_header or _contains_marker(col_header)) and table is not None:
        col_header = _secondary_col_header(table, row_index, col_index) or col_header

    # ── 日期格被非日期列头覆盖 ──
    # 签字/记录/标准块行的纯日期格（____年__月__日）常落在标准号/备注/是否归档等列头下，
    # 左侧行标签才是真实含义（检测日期/标准块有效期/日 期）。按行标签命名并清列头，
    # 避免出现「YY0621.1-2016」「备注」等张冠李戴的字段名。
    if (
        _is_date_cell(template_text)
        and row_label
        and any(w in row_label for w in ("日期", "时间", "有效期"))
        and not any(w in (col_header or "") for w in ("日期", "时间", "有效期"))
    ):
        rl = row_label.strip().replace(" ", "")
        if rl in ("日期", "时间", "有效期"):
            # 通用日期标签：补上左侧最近的签字角色/对象名（如「检测人员 日期」）
            prefix_name = ""
            for c_col, _, c_text in row_cells:
                if c_col >= col_index:
                    break
                if c_text and not _contains_marker(c_text) and not any(w in c_text for w in ("日期", "时间", "有效期")):
                    prefix_name = c_text.strip()
            label = f"{prefix_name} {rl}" if prefix_name else rl
        else:
            label = rl
        row_label = label
        col_header = ""
        return label, row_label, col_header

    # ── 签字/角色行 ──
    # 角色行（检测人员/核验人员/技术负责人…）的空白/勾选格改按「角色 + 行内标签」命名，
    # 避免显示成「检测人员 · 编号或文件名」「签字 · 备注」这类错位字段名。
    role = _nearest_role(row_cells, col_index) or _role_row_name(row_cells)
    if role:
        inline = row_label.replace(" ", "").strip()
        # 有签名相关列头（姓名/签字/日期/备注/结论…）时优先尊重列头，
        # 避免独立签字表里「备注」空列被误当姓名格导致「检测人员 姓名」重复。
        if (
            col_header
            and not _contains_marker(col_header)
            and any(w in col_header for w in ("姓名", "签字", "签名", "日期", "时间", "备注", "结论", "判定", "结果"))
        ):
            ch = col_header.strip()
            label = f"{role} {ch}" if ch not in role else ch
        elif _is_date_cell(template_text):
            # 角色行的纯日期格（____年__月__日）紧邻角色名、没有独立的「日期」标签格，
            # 此时 row_label 是角色名而非日期词，走不到上面的日期修复分支，这里单独补日期名。
            label = f"{role} 日期"
        elif "签名" in template_text or "签字" in template_text:
            # 复合格「签名：____ 日期：____年__月__日」整体按签字格命名。
            label = f"{role} 签字"
        elif inline in ("签字", "签名"):
            label = f"{role} 签字"
        elif _is_first_blank_after_role(row_cells, col_index):
            # 角色名后的第一个空白格 = 姓名格。
            label = f"{role} 姓名"
        elif inline == role.replace(" ", ""):
            # 后续空白格：右侧紧邻「签字/签名/日期」标签则按之命名，
            # 避免「检测人员 | 姓名 | 签字：」出现两个「检测人员 姓名」。
            right = _right_label(row_cells, col_index)
            if any(w in right for w in ("签字", "签名")):
                label = f"{role} 签字"
            elif any(w in right for w in ("日期", "时间", "有效期")):
                label = f"{role} 日期"
            else:
                label = f"{role} 姓名"
        elif inline:
            label = f"{role} {row_label.strip()}"
        else:
            label = role
        row_label = label
        col_header = ""
        return label, row_label, col_header

    meaningful_template = re.sub(r"[_＿]+", "", template_text).strip(" /：:；;")
    prefix = ""
    if template_text and meaningful_template and not template_text.startswith("□"):
        prefix = re.split(r"_{2,}|＿{2,}|□", template_text, maxsplit=1)[0].strip(" /：:；;")

    if not template_text and immediate_left and not _contains_marker(immediate_left):
        label = immediate_left
    elif col_header and not _contains_marker(col_header):
        label = col_header
    elif prefix:
        label = prefix
    elif row_label and not _contains_marker(row_label):
        label = row_label
    else:
        label = f"表{table_index + 1}第{row_index + 1}行第{col_index + 1}列"

    if row_index > 0 and col_header and label == col_header:
        label = f"{label}（第{row_index}条）"

    # 结果列（确认/判定/结果/结论等通用词）及签字表「日期」在多行重复时，
    # 带上该行行首确认项/签字人角色名以区分，避免母版过程确认里出现多个同名
    # 「判定」「确认」「日期」无法区分。
    _RESULT_WORDS = {"确认", "判定", "结果", "结论", "日期"}
    if label in _RESULT_WORDS:
        # 取左侧最近的、非通用词的角色标签（如「检测人员」「核验人员」），
        # 而非最左的，以支持同一行横排多个签字角色（检测人员…日期 核验人员…日期）。
        first_label = ""
        for candidate_col, _, candidate_text in row_cells:
            if candidate_col >= col_index:
                break
            if (
                candidate_text
                and not _contains_marker(candidate_text)
                and candidate_text.strip() not in _RESULT_WORDS
            ):
                first_label = candidate_text
        if first_label and first_label != label:
            label = f"{first_label} {label}"
            row_label = label

    return label, row_label, col_header


# ── 模板清单 ──────────────────────────────────────────────

def template_manifest(template_path: Path | str) -> list[dict[str, Any]]:
    """解析 DOCX 模板，返回所有可填充字段的清单"""
    path = Path(template_path)
    if not path.exists():
        return []
    doc = Document(str(path))
    sections = _body_table_sections(doc)
    fields: list[dict[str, Any]] = []
    for table_index, table in enumerate(doc.tables):
        headers = _table_headers(table)
        seen_cells = set()
        for row_index, row in enumerate(table.rows):
            row_cells = _unique_row_cells(row)
            for col_index, cell, template_text in row_cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                if row_index == 0 and not _contains_marker(template_text):
                    continue
                if not _contains_marker(template_text):
                    continue
                label, row_label, col_header = _field_label(
                    table_index, row_index, col_index, row_cells, headers, template_text, table
                )
                fields.append({
                    "key": f"t{table_index}_r{row_index}_c{col_index}",
                    "section": sections[table_index] if table_index < len(sections) else f"表{table_index + 1}",
                    "table": table_index, "row": row_index, "col": col_index,
                    "label": label, "row_label": row_label, "col_header": col_header,
                    "template_text": template_text,
                    "input_type": _infer_input_type(template_text),
                    "position": f"表{table_index + 1}-R{row_index + 1}C{col_index + 1}",
                })
    return fields


# ── 单元格填入 ────────────────────────────────────────────

def _clone_rpr(source_run, target_run):
    """Copy run properties from source to target."""
    if source_run is None:
        return
    src_rpr = source_run._r.find(qn("w:rPr"))
    if src_rpr is None:
        return
    tgt_rpr = target_run._r.find(qn("w:rPr"))
    if tgt_rpr is None:
        from lxml import etree
        tgt_rpr = etree.SubElement(target_run._r, qn("w:rPr"))
    for child in list(src_rpr):
        tgt_rpr.append(deepcopy(child))


def _write_cell_text(cell, original: str, value: Any, changed: bool = False) -> None:
    text = "" if value is None else str(value)
    paragraphs = list(cell.paragraphs)
    paragraph = paragraphs[0] if paragraphs else cell.add_paragraph()
    source_run = next((r for p in paragraphs for r in p.runs), None)
    runs = [run for p in paragraphs for run in p.runs]
    if not runs:
        runs = [paragraph.add_run("")]
        _clone_rpr(source_run, runs[0])
    runs[0].text = text
    runs[0].font.color.rgb = RED if changed else BLACK
    for run in runs[1:]:
        run.text = ""
        run.font.color.rgb = BLACK


def fill_exact_template(
    template_path: Path | str,
    values: dict[str, Any],
    changed_keys: set[str] | None = None,
) -> _Document:
    """将 values 精确填入模板对应单元格，不增删任何表格结构"""
    path = Path(template_path)
    if not path.exists():
        raise FileNotFoundError(f"受控原始记录模板不存在：{template_path}")
    changed_keys = changed_keys or set()
    doc = Document(str(path))
    manifest_map = {field["key"]: field for field in template_manifest(path)}
    for key, value in values.items():
        field = manifest_map.get(key)
        if not field:
            continue
        table = doc.tables[field["table"]]
        row = table.rows[field["row"]]
        if field["col"] >= len(row.cells):
            continue
        cell = row.cells[field["col"]]
        _write_cell_text(cell, str(field.get("template_text", "") or ""), value, key in changed_keys)
    return doc


# ── 电子签名 ──────────────────────────────────────────────

def _signature_path(signature_dir: Path, username: str) -> Path | None:
    """查找用户签名图片"""
    if not username or not signature_dir.exists():
        return None
    for ext in (".png", ".jpg", ".jpeg"):
        candidate = signature_dir / f"{username}{ext}"
        if candidate.exists():
            return candidate
    return None


def _put_signature(cell, signature_dir: Path, username: str, date_text: str = ""):
    """在单元格中放置签名图片"""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    path = _signature_path(signature_dir, username)
    if path:
        paragraph.add_run().add_picture(str(path), width=Inches(0.82))
    else:
        paragraph.add_run("【未配置签名图片】")
    if date_text:
        paragraph.add_run(f"  {_china_date_str(date_text)}")


def _apply_record_signatures(
    doc: _Document, record: dict, task: dict | None,
    signature_dir: Path,
):
    """在 DOCX 中填入实验员/复核员的电子签名"""
    tester = record.get("owner") or (task.get("assignee") if task else "")
    reviewer = task.get("reviewer") if task else ""
    tester_date = record.get("tester_signed_at") or record.get("updated_at") or ""
    reviewer_date = record.get("reviewer_signed_at") or ""

    tester_tokens = ("检测人员", "实验员")
    reviewer_tokens = ("核验人员", "复核人员", "核验员", "复核员")

    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        observer_columns = [i for i, text in enumerate(headers) if "观察者签字" in text]
        for row in table.rows[1:]:
            for col in observer_columns:
                if col < len(row.cells) and row.cells[col].text.strip() not in ("", "/", "不适用"):
                    _put_signature(row.cells[col], signature_dir, tester, tester_date)

        for row_index, row in enumerate(table.rows):
            cells = []
            seen = set()
            for index, cell in enumerate(row.cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                cells.append((index, cell, cell.text.strip()))
            row_text = " ".join(text for _, _, text in cells)
            exact_role_row = any(
                text.strip(" ：:（）()") in tester_tokens + reviewer_tokens
                for _, _, text in cells
            )
            signature_row = (
                any(token in row_text for token in ("签字", "签名", "日期", "年__", "年__月", "/年/月"))
                or (row_index > 0 and exact_role_row and len(cells) <= 8)
            )
            if not signature_row:
                continue
            for position, (_index, _cell, label) in enumerate(cells):
                username = ""
                signed_at = ""
                if any(token in label for token in tester_tokens):
                    username, signed_at = tester, tester_date
                elif any(token in label for token in reviewer_tokens):
                    username, signed_at = reviewer, reviewer_date
                if not username or position + 1 >= len(cells):
                    continue
                target_position = position + 1
                for later in range(position + 1, len(cells) - 1):
                    if cells[later][2].strip("：: ") in ("签字", "签名"):
                        target_position = later + 1
                        break
                _put_signature(cells[target_position][1], signature_dir, username, signed_at)

            # 合并确认格（如 R004）
            for _index, cell, text in cells:
                if "确认人" in text and "复核" in text:
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.add_run("确认：")
                    tester_path = _signature_path(signature_dir, tester)
                    if tester_path:
                        p.add_run().add_picture(str(tester_path), width=Inches(0.62))
                    p.add_run("  复核：")
                    reviewer_path = _signature_path(signature_dir, reviewer)
                    if reviewer_path:
                        p.add_run().add_picture(str(reviewer_path), width=Inches(0.62))


# ── 回退模板 ──────────────────────────────────────────────

def _fallback_docx(record: dict) -> BytesIO:
    """没有受控模板时的兜底 DOCX"""
    payload = record.get("payload", {})
    if isinstance(payload, str):
        import json
        try:
            payload = json.loads(payload)
        except Exception:
            payload = {}
    doc = Document()
    doc.add_heading(record.get("experiment") or "实验原始记录", 0)
    doc.add_paragraph("当前实验尚未配置受控原始记录模板。")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "记录值"
    template_fields = payload.get("template_fields") or payload.get("_template_fields") or {}
    if isinstance(template_fields, list):
        for tf in template_fields:
            cells = table.add_row().cells
            cells[0].text = str(tf.get("key") or tf.get("label") or "")
            cells[1].text = str(tf.get("value") or "")
    else:
        for key, value in template_fields.items():
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = BLACK
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ── 导出入口 ──────────────────────────────────────────────

# DB kind → Mapper kind
_KIND_TO_MAPPER = {
    "roughness": "rough",
    "crack": "mc_crack",
    "xray": "xray",
    "warpage": "warp",
    "cte": "cte",
    "thermal_shock": "shock",
    "bending": "bend",
    "vickers": "hv",
    "thickness": "thickness",
    "color_stability": "color",
    "fixed_denture": "fixed_denture",
    "removable_denture": "removable_denture",
    "density": "density",
    "tarnish": "tarnish",
}

# experiment_code → mapper kind 的直接映射（用于 DB 中无 kind 字段的回退）
_CODE_TO_MAPPER = {
    "I001": "rough", "I002": "mc_crack", "I003": "xray",
    "I004": "warp", "I005": "cte", "I006": "shock",
    "I007": "bend", "I008": "hv", "I009": "thickness",
    "I010": "color", "I011": "fixed_denture", "I012": "removable_denture",
    "I013": "density", "I014": "tarnish",
}


def _get_kind(record: dict, task: dict | None) -> str:
    """从 record/task 提取 mapper kind 值（多级回退）"""
    kind = record.get("kind") or (task.get("kind") if task else None) or ""
    if kind:
        mapped = _KIND_TO_MAPPER.get(kind)
        if mapped:
            return mapped
    # 通过 experiment_code 回退
    exp_code = record.get("experiment_code") or (task.get("experiment_code") if task else None) or ""
    if exp_code:
        mapped = _CODE_TO_MAPPER.get(exp_code)
        if mapped:
            return mapped
    return kind


def _commission_header_values(
    manifest: list[dict[str, Any]],
    commission: dict | None,
    groups: list | None,
    samples: list | None,
    task: dict | None,
    record: dict,
) -> dict[str, str]:
    """根据委托/样品组/样品/任务数据，生成原始记录表头部“样品与委托信息”各单元格的值。

    采用 label 归一化匹配（文本字段）+ checkbox 内容匹配（检验类别/接收状态），
    适用于所有受控记录模板，不依赖具体表格索引。
    """
    commission = commission or {}
    groups = groups or []
    samples = samples or []
    task = task or {}

    # 该任务对应的样品组
    group_no = task.get("group_no") or ""
    sg: dict = {}
    for g in groups:
        if g.get("group_no") == group_no:
            sg = g
            break
    if not sg and groups:
        sg = groups[0]

    sample_nos = [str(s.get("sample_no", "")) for s in samples if s.get("sample_no")]
    if not sample_nos:
        raw_nos = str(task.get("sample_nos", "") or "").replace("，", ",")
        sample_nos = [x.strip() for x in raw_nos.split(",") if x.strip()]

    quantity = sg.get("quantity") or len(sample_nos) or ""
    production_unit = sg.get("production_org_name") or commission.get("production_org_name") or ""

    test_date = ""
    for src in (task.get("experiment_started_at"), task.get("experiment_ended_at"),
                record.get("tester_signed_at"), record.get("created_at")):
        if src:
            test_date = _china_date_str(src)
            break
    receive_date = str(commission.get("commission_date") or "")[:10]

    lab_sample_nos = "、".join(sample_nos)
    client_lot = sg.get("product_no") or sg.get("batch_no") or ""

    # 归一化字段名（精确匹配，避免“样品编号”误匹配“样品编号/批号”）
    text_map: dict[str, str] = {
        "委托单位": commission.get("client_name") or "",
        "委托单位送检单位": commission.get("client_name") or "",
        "委托单位地址": commission.get("client_address") or "",
        "委托单位编号": commission.get("commission_no") or "",
        "委托编号": commission.get("commission_no") or "",
        "生产单位": production_unit,
        "生产厂家": production_unit,
        "样品名称": sg.get("sample_name") or "",
        "样品编号批号": client_lot,
        "产品编号批号": client_lot,
        "样品编号": lab_sample_nos,
        "实验室样品编号": lab_sample_nos,
        "样品数量": str(quantity),
        "样品规格": sg.get("model") or "",
        "样品规格型号": sg.get("model") or "",
        "规格型号结构功能": sg.get("model") or "",
        "规格型号功能分类": sg.get("model") or "",
        "材料工艺": sg.get("material_name") or "",
        "基托支架材料": sg.get("material_name") or "",
        "接收日期": receive_date,
        "检测日期": test_date,
        "检测地点": task.get("detection_location") or "",
        "原始记录编号": record.get("record_no") or task.get("task_no") or "",
        "记录编号": record.get("record_no") or task.get("task_no") or "",
        "产品生产日期": sg.get("production_date") or "",
        "样品生产日期": sg.get("production_date") or "",
    }

    values: dict[str, str] = {}
    for field in manifest:
        key = field["key"]
        original = str(field.get("template_text", "") or "")
        is_checkbox = "□" in original or "☐" in original

        # 1) 文本字段：按 row_label → label 归一化精确匹配
        if not is_checkbox:
            for label_src in (field.get("row_label"), field.get("label"), field.get("col_header")):
                norm = _norm_label(label_src)
                if not norm:
                    continue
                if norm in text_map and text_map[norm]:
                    values[key] = _compose_cell_text(original, text_map[norm])
                    break
            continue

        # 2) checkbox 内容匹配（仅处理与委托相关的通用勾选）
        if "完好" in original and any(w in original for w in ("异常", "污染", "破损", "缺件")):
            values[key] = _select_checkbox(original, sg.get("condition") or "完好")
            continue
        if ("委托检测" in original or "委托检验" in original) and "型式检验" in original:
            preferred = "委托检测" if "委托检测" in original else "委托检验"
            values[key] = _select_checkbox(original, preferred)
            continue
        if "委托方提供" in original:
            values[key] = _select_checkbox(original, "委托方提供")
            continue

        # 3) 头部 checkbox：规格型号 / 材料工艺 —— 对齐参考仓库 business_record_engine
        #    由样品组权威数据勾选并回填，避免实验员在⑤里看到空的规格型号/材料勾选框。
        label_norm = _norm_label(
            " ".join([
                field.get("label", ""), field.get("row_label", ""),
                field.get("col_header", ""), original,
            ])
        )
        if any(t in label_norm for t in ("样品规格", "型号规格", "规格型号")):
            model = str(sg.get("model") or "").strip()
            if model and model != "-":
                values[key] = _select_checkbox_fill(original, model)
                continue
        if any(t in label_norm for t in ("材料工艺", "材料名称")):
            material = str(sg.get("material_name") or "").strip()
            if material and material != "-":
                values[key] = _select_checkbox_fill(original, material)
                continue

    return values


def compute_template_values(
    record: dict,
    task: dict | None,
    template_dir: Path,
    commission: dict | None = None,
    groups: list | None = None,
    samples: list | None = None,
    registry: dict | None = None,
    finalize: bool = True,
    display_names: dict[str, str] | None = None,
) -> tuple[str | None, Path | None, dict[str, Any]]:
    """计算受控记录模板的完整填充值（后端预填 + 实验员 _template_fields）。

    返回 (template_name, template_path, template_fields)。找不到模板时返回 (None, None, {})。
    该函数是「导出 DOCX」与「⑤母版过程确认」的统一数据源，保证两者逐格一致。
    """
    payload = record.get("payload", {})
    if isinstance(payload, str):
        import json
        try:
            payload = json.loads(payload)
        except Exception:
            payload = {}

    # 确定模板文件
    template_name = record.get("record_template_file") or payload.get("_template_name") or ""
    if not template_name:
        # 尝试从实验代码推断 template_code
        experiment_code = task.get("experiment_code") if task else record.get("experiment_code", "")
        # Resolve kind → template code via the same mapping used in experiment_config
        kind = (task.get("kind") if task else None) or record.get("kind", "")
        if not kind and experiment_code:
            kind = experiment_code  # fallback

        # Map kind / experiment_code to template code
        _KIND_MAP = {
            "rough": "R001", "mc_crack": "R004", "xray": "R005",
            "warp": "R006", "cte": "R007", "shock": "R009",
            "bend": "R010", "hv": "R011", "color": "R012", "thickness": "R013",
            "fixed_denture": "R014", "removable_denture": "R015",
            "density": "R016", "tarnish": "R017",
        }
        _CODE_TO_TEMPLATE = {
            "I001": "R001", "I002": "R004", "I003": "R005",
            "I004": "R006", "I005": "R007", "I006": "R009",
            "I007": "R010", "I008": "R011", "I009": "R013",
            "I010": "R012", "I011": "R014", "I012": "R015",
            "I013": "R016", "I014": "R017",
        }
        template_code = _KIND_MAP.get(kind) or _CODE_TO_TEMPLATE.get(experiment_code, experiment_code)

        # Search templates: try R001_*.docx, then RECORD_R001_*.docx, then SOP_R001_*.docx
        if template_dir.exists() and template_code:
            for prefix in (template_code + "_", template_code + ".", "RECORD_" + template_code, "SOP_" + template_code):
                for f in template_dir.iterdir():
                    if f.suffix != '.docx':
                        continue
                    if f.name.startswith(prefix):
                        template_name = f.name
                        break
                if template_name:
                    break

    template_path = template_dir / template_name if template_name else None
    if not template_name or not template_path or not template_path.exists():
        return None, None, {}

    # 提取模板字段值（兼容 _template_fields 和 template_fields 两种键名）
    template_fields = payload.get("template_fields") or payload.get("_template_fields") or {}
    if isinstance(template_fields, list):
        template_fields = {tf.get("key"): tf.get("value") for tf in template_fields}

    # 始终尝试通过受控映射填充空值
    try:
        from app.services.controlled_template_mappings import apply_controlled_mapping
        kind = _get_kind(record, task)
        if kind and template_name:
            template_path_full = template_dir / template_name
            dn = display_names or {}
            def _name(u: str | None) -> str:
                u = u or ""
                return dn.get(u, u)
            context = {
                "experiment": record.get("experiment") or (task.get("experiment") if task else ""),
                "experiment_code": record.get("experiment_code") or (task.get("experiment_code") if task else ""),
                "tester": _name(record.get("owner", "")),
                "operator": _name(record.get("owner") or (task.get("assignee") if task else "")),
                "reviewer": _name((task.get("reviewer") if task else "") or record.get("reviewer", "")),
            }
            # 构建 mapper 期望的 business_record 结构（parameters + rows）
            business_record = dict(payload)
            business_record.setdefault("parameters", payload.get("_form") or {})
            business_record.setdefault("rows", payload.get("_rows") or [])
            # 兼容记录顶层的游离测量字段（如 _standard_block_measured）注入 parameters，
            # 使受控映射能读取实验员实际记录的实测值。
            for extra_key in ("_standard_block_measured", "standard_block_measured"):
                extra_val = payload.get(extra_key)
                if extra_val not in (None, ""):
                    business_record["parameters"].setdefault(extra_key, extra_val)
            mapped_fields, touched_keys = apply_controlled_mapping(
                str(template_path_full), kind, {}, context, business_record, "",
                registry=registry, finalize=finalize,
            )
            if mapped_fields:
                # 合并策略：受控映射（由 _form/_rows/_equipment_checks 等实验员数据生成）优先；
                # 历史 _template_fields 渲染缓存仅用于补充映射未显式生成的非勾选单元格。
                merged = dict(mapped_fields)
                existing = dict(template_fields) if template_fields else {}
                for k, v in existing.items():
                    # 受控映射已显式生成的单元格，以实验员数据/SOP 为准
                    if k in touched_keys:
                        continue
                    existing_str = str(v or "")
                    # 空值/未填充标记 → 不保留
                    if not existing_str.strip() or "____" in existing_str:
                        continue
                    # 勾选值（□/☐/☑）由受控映射与委托头部权威生成，不保留历史渲染结果
                    if "□" in existing_str or "☐" in existing_str or "☑" in existing_str:
                        continue
                    merged[k] = v
                template_fields = merged
    except Exception as _e:
        import traceback
        traceback.print_exc()

    # 填入委托/样品头部信息（委托单位、生产单位、样品编号、接收日期等）
    try:
        header_values = _commission_header_values(
            template_manifest(template_path), commission, groups, samples, task, record
        )
        if header_values:
            merged = dict(template_fields or {})
            for k, v in header_values.items():
                merged[k] = v
            template_fields = merged
    except Exception:
        import traceback
        traceback.print_exc()

    return template_name, template_path, template_fields


def template_supplement_requirements(
    template_path: Path,
    values: dict[str, Any],
) -> list[dict[str, Any]]:
    """返回后端预填后仍未完成的模板字段清单（⑤母版过程确认需实验员补充的部分）。

    对齐参考仓库 business_record_engine.template_supplement_requirements：
    - 勾选框未选中（无 ☑）视为未完成；
    - 文本字段仍含空白标记（___/＿/…）视为未完成；
    - "/"、"不适用" 视为已处理，跳过。
    """
    requirements: list[dict[str, Any]] = []
    for field in template_manifest(template_path):
        original = str(field.get("template_text", "") or "")
        current = str(values.get(field["key"], "") or "").strip()
        if current in {"/", "不适用"}:
            continue
        is_checkbox = "□" in original or "☐" in original
        if is_checkbox:
            # 勾选框：未被选中（无 ☑）视为未完成
            if "☑" not in current:
                requirements.append(field)
        else:
            # 文本字段：仅在模板本身含空白占位（___/＿/…）且未被填满时视为未完成。
            # 纯标签单元格（template_text 无占位符）不在此列。
            effective = current or original
            if BLANK_RE.search(effective):
                requirements.append(field)
    return requirements


def export_record_docx(
    record: dict,
    task: dict | None,
    template_dir: Path,
    signature_dir: Path,
    commission: dict | None = None,
    groups: list | None = None,
    samples: list | None = None,
    registry: dict | None = None,
    display_names: dict[str, str] | None = None,
) -> bytes:
    """生成填入实验数据 + 电子签名的受控 DOCX 字节流

    record: 记录字典（含 payload）
    task:   任务字典（含 reviewer, assignee 等）
    registry: 受控模板映射的 DB 权威值（{"db_mappings": [...], "extra_json": {...}}），
              由异步调用方从 experiment_config.load_mapping_registry 读入；无则硬编码兜底。
    """
    template_name, template_path, template_fields = compute_template_values(
        record, task, template_dir, commission, groups, samples, registry,
        display_names=display_names,
    )
    if template_path is None:
        return _fallback_docx(record).getvalue()

    # 收集变化字段（用于红色标记）
    changed_keys = set()
    if int(record.get("version", 1) or 1) > 1:
        changes = record.get("changes") or []
        for item in changes:
            field_name = str(item.get("field_name", ""))
            if field_name.startswith("template_fields."):
                changed_keys.add(field_name.split("template_fields.", 1)[1])

    # 填入模板
    doc = fill_exact_template(template_path, template_fields, changed_keys)

    # 放置电子签名
    _apply_record_signatures(doc, record, task, signature_dir)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
