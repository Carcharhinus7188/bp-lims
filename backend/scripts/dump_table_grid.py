# -*- coding: utf-8 -*-
"""打印指定模板中某个表的完整网格（所有单元格，含不可填），用于坐标对齐审计。

用法:
  cd backend && .venv/Scripts/python.exe scripts/dump_table_grid.py <template_path> <table_index> [table_index...]
"""
from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from docx import Document
from app.services.record_word_engine import template_manifest


def _clean(s):
    return "".join(s.split())


def dump(template_path: str, table_indices: list[int]) -> None:
    doc = Document(template_path)
    manifest = template_manifest(template_path)
    fillable = {(f["table"], f["row"], f["col"]) for f in manifest}

    for ti in table_indices:
        if ti >= len(doc.tables):
            print(f"[表{ti}] 不存在（模板共 {len(doc.tables)} 张表）\n")
            continue
        table = doc.tables[ti]
        print(f"\n{'='*100}\n表{ti}（共 {len(table.rows)} 行 × {len(table.columns)} 列）\n{'='*100}")
        # 去重合并单元格：记录每个 tc 首次出现位置
        seen_tc = set()
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                if cell._tc in seen_tc:
                    cells.append((ci, None, ""))  # 合并单元格（水平/垂直延续）
                else:
                    seen_tc.add(cell._tc)
                    cells.append((ci, cell, _clean(cell.text)))
            print(f"--- 行 {ri} ---")
            for ci, cell, text in cells:
                if cell is None:
                    continue
                mark = "FILL" if (ti, ri, ci) in fillable else "    "
                print(f"  [{mark}] r{ri} c{ci}: {text!r}")
        print()


if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) < 2:
        print("用法: dump_table_grid.py <template_path> <table_index> [table_index...]")
        sys.exit(1)
    tp = args[0]
    tis = [int(x) for x in args[1:]]
    dump(tp, tis)
