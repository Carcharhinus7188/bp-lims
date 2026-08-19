"""根据《试样编码清单》Excel 全删重建样品资料库 sample_catalog。

数据源：<repo>/2026.8.18试样编码清单_GDHD按材料顺序重新编码版.xlsx 的「重新编码清单」sheet
（8 列：试样编码 / 试样名称 / 规格型号 / 主要材料 / 工艺 / 单位 / 检验项目 / 检验依据）

规则：
- 全删重建：删除 sample_code 形如 SY-*/GD-*/HD-*/QT-* 的试样记录，以及 sample_code 为空的废弃记录；
  保留牙科材料 SL* 记录（不在本文件内）。
- 字段映射：检验项目 → detection_method，检验依据 → detection_basis，其余按列名一一对应。
- 类别派生：SY→实验试样，GD→定制式固定义齿，HD→定制式活动义齿，QT→其他/待确认。
- 连接检测实验 experiment_codes：
    固定义齿 → I011（金属/金属-陶瓷 +I001）；活动义齿 → I012；
    弯曲→I007；抗晦暗→I014；金瓷结合→I002；粗糙度+硬度+密度→I001+I008+I013；
    线胀系数→I005；尺寸/翘曲→I009+I004；拉伸/耐腐蚀/弹性模量→[]（暂不关联）。
- 已建委托的 sample_groups.catalog_id 按 sample_code 重映射到新 id。

用法：cd backend && .venv/Scripts/python.exe -X utf8 scripts/rewrite_sample_catalog.py [--apply]
默认 dry-run；加 --apply 才写库。
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import asyncpg
import openpyxl

BACKEND = Path(__file__).parent.parent
XLSX = BACKEND.parent / "2026.8.18试样编码清单_GDHD按材料顺序重新编码版.xlsx"

DB = dict(
    host="localhost",
    user="postgres",
    password="123456",
    database="bplab",
)

METAL_TOKENS = ["纯钛", "钴铬", "镍铬", "钛合金", "金合金"]


def clean(s) -> str:
    if s is None:
        return ""
    s = str(s).strip()
    s = s.replace("　", " ")  # 全角空格
    s = re.sub(r"\s+", "", s)     # 去掉换行与空白（依据原文用「；」分隔）
    return s.strip("；; ")


def category_for(code: str) -> str:
    if not code:
        return "其他/待确认"
    p = code.split("-")[0].upper()
    return {
        "SY": "实验试样",
        "GD": "定制式固定义齿",
        "HD": "定制式活动义齿",
        "QT": "其他/待确认",
    }.get(p, "其他/待确认")


def material_suffix_for(code: str) -> str | None:
    if not code:
        return None
    parts = code.split("-")
    if len(parts) == 4 and parts[0].upper() == "SY":
        return parts[3]
    return None


def propose(name: str, material: str) -> list[str]:
    name = name or ""
    mat = material or ""
    if name == "定制式固定义齿":
        codes = ["I011"]
        if any(t in mat for t in METAL_TOKENS):
            codes.append("I001")
        return codes
    if name == "定制式活动义齿":
        return ["I012"]
    if "试样" in name:
        if "弯曲" in name:
            return ["I007"]
        if "抗晦暗" in name:
            return ["I014"]
        if "金瓷结合" in name or "结合强度" in name:
            return ["I002"]
        if "粗糙度" in name and "硬度" in name and "密度" in name:
            return ["I001", "I008", "I013"]
        if "线胀系数" in name or "热膨胀" in name:
            return ["I005"]
        if "尺寸" in name or "翘" in name:
            return ["I009", "I004"]
        return []
    return []


def load_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["重新编码清单"]

    # 定位表头行（A 列 = 试样编码）
    header_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value and str(ws.cell(r, 1).value).strip() == "试样编码":
            header_row = r
            break
    if header_row is None:
        raise RuntimeError("未找到「试样编码」表头行")

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        code = clean(ws.cell(r, 1).value)
        if not code:
            continue
        name = clean(ws.cell(r, 2).value)
        model = clean(ws.cell(r, 3).value) or "标准"
        material = clean(ws.cell(r, 4).value)
        process = clean(ws.cell(r, 5).value)
        unit = clean(ws.cell(r, 6).value) or "件"
        method = clean(ws.cell(r, 7).value)
        basis = clean(ws.cell(r, 8).value)
        codes = sorted(set(propose(name, material)))
        rows.append({
            "sample_code": code,
            "sample_name": name,
            "model": model,
            "material_name": material,
            "process": process or None,
            "unit": unit,
            "category": category_for(code),
            "detection_method": method or None,
            "detection_basis": basis or None,
            "experiment_codes": codes,
            "material_suffix": material_suffix_for(code),
        })
    return rows


async def main(apply: bool) -> None:
    rows = load_rows()
    print(f"从 Excel 解析到 {len(rows)} 条试样记录。")

    conn = await asyncpg.connect(**DB)
    try:
        # 快照：现有 sample_groups.catalog_id → sample_code（重建后重映射用）
        snap = await conn.fetch(
            "SELECT sg.id, sg.catalog_id, sc.sample_code "
            "FROM sample_groups sg LEFT JOIN sample_catalog sc ON sc.id = sg.catalog_id "
            "WHERE sg.catalog_id IS NOT NULL"
        )

        # 统计当前状态
        before_specimen = await conn.fetchval(
            "SELECT count(*) FROM sample_catalog WHERE sample_code ~ '^(SY|GD|HD|QT)-'"
        )
        before_sl = await conn.fetchval(
            "SELECT count(*) FROM sample_catalog WHERE sample_code LIKE 'SL%'"
        )
        before_null = await conn.fetchval(
            "SELECT count(*) FROM sample_catalog WHERE sample_code IS NULL"
        )
        print(f"当前库：试样 {before_specimen} 条，牙科材料 SL {before_sl} 条，无编码废弃 {before_null} 条。")

        # 待删除：试样 + 无编码废弃（保留 SL）
        to_delete = await conn.fetchval(
            "SELECT count(*) FROM sample_catalog WHERE sample_code ~ '^(SY|GD|HD|QT)-' OR sample_code IS NULL"
        )
        print(f"计划删除 {to_delete} 条（试样 + 无编码废弃），保留 SL 牙科材料。")

        # 预览差异
        name_updates = {
            "SY-CC": ("尺寸、外观检查试样", "尺寸、翘起变形试样"),
            "SY-JC": ("金瓷结合强度试样", "金瓷结合性能试样"),
        }
        print("\n预览：名称更正")
        for k, (old, new) in name_updates.items():
            print(f"  {k}* : {old} → {new}")

        if not apply:
            print("\n（dry-run，未写入。确认后加 --apply 执行。）")
            print("首批 3 条示例：")
            for r in rows[:3]:
                print("  ", json.dumps(r, ensure_ascii=False))
            return

        # 事务：删 → 插 → 重映射
        async with conn.transaction():
            await conn.execute(
                "DELETE FROM sample_catalog WHERE sample_code ~ '^(SY|GD|HD|QT)-' OR sample_code IS NULL"
            )
            new_ids = {}
            for r in rows:
                new_id = await conn.fetchval(
                    """
                    INSERT INTO sample_catalog
                      (sample_code, sample_name, model, material_name, process, unit,
                       category, detection_method, detection_basis, experiment_codes,
                       material_suffix, enabled)
                    VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,TRUE)
                    RETURNING id
                    """,
                    r["sample_code"], r["sample_name"], r["model"], r["material_name"],
                    r["process"], r["unit"], r["category"], r["detection_method"],
                    r["detection_basis"], json.dumps(r["experiment_codes"]),
                    r["material_suffix"],
                )
                new_ids[r["sample_code"]] = new_id

            # 重映射已建委托的 catalog_id
            remapped = 0
            for sg_id, old_cid, old_code in snap:
                if old_code in new_ids:
                    await conn.execute(
                        "UPDATE sample_groups SET catalog_id=$1 WHERE id=$2",
                        new_ids[old_code], sg_id,
                    )
                    remapped += 1

        after = await conn.fetchval("SELECT count(*) FROM sample_catalog")
        print(f"\n✓ 完成：删除 {to_delete} 条，插入 {len(rows)} 条，重映射 catalog_id {remapped} 条。")
        print(f"  库内现共 {after} 条（含 SL 牙科材料）。")

    finally:
        await conn.close()


if __name__ == "__main__":
    import asyncio
    asyncio.run(main("--apply" in sys.argv))
